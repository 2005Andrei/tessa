import io
from docx import Document
from docx.shared import Inches

def create_doc():
    doc = Document()
    doc.add_heading("A heeading", 0)
    doc.add_paragraph("A paragraph")
    
    return doc
